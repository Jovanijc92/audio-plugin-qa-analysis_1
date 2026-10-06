"""
Build the CSV reports and the Excel workbook from the cleaned database.

Run sql/03_data_cleaning.sql before this script. The workbook formulas
recalculate the summary numbers from the cleaned test and bug rows, so the
Excel file can be explained without treating Python as the analysis tool.
"""

import csv
import sqlite3
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
DB_PATH = ROOT / "database" / "audio_plugin_qa.db"
EXCEL_PATH = ROOT / "excel" / "audio_plugin_qa_analysis.xlsx"

NAVY = "1F4E79"
LIGHT = "E7EEF5"
PALE = "F7F9FC"
ORANGE = "FCE4D6"
LINE = "D9E2F0"
WHITE = "FFFFFF"

title_font = Font(name="Calibri", size=18, bold=True, color=NAVY)
section_font = Font(name="Calibri", size=13, bold=True, color=NAVY)
header_font = Font(name="Calibri", size=11, bold=True, color=WHITE)
label_font = Font(name="Calibri", size=11, color="333333")
note_font = Font(name="Calibri", size=10, italic=True, color="555555")
total_font = Font(name="Calibri", size=11, bold=True, color="333333")

header_fill = PatternFill("solid", fgColor=NAVY)
light_fill = PatternFill("solid", fgColor=LIGHT)
pale_fill = PatternFill("solid", fgColor=PALE)
total_fill = PatternFill("solid", fgColor="D6E3F0")
alert_fill = PatternFill("solid", fgColor=ORANGE)

thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
center = Alignment(horizontal="center", vertical="center")
left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)


def split_statements(sql_text):
    statements = []
    current = []
    for line in sql_text.splitlines():
        current.append(line)
        if line.strip().endswith(";"):
            statements.append("\n".join(current).strip())
            current = []
    tail = "\n".join(current).strip()
    if tail:
        statements.append(tail)
    return statements


def statement_body(statement):
    lines = []
    for line in statement.splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("--"):
            continue
        lines.append(stripped)
    return " ".join(lines).strip()


def export_target(statement):
    for line in statement.splitlines():
        stripped = line.strip()
        if stripped.startswith("-- EXPORT:"):
            return stripped.split(":", 1)[1].strip()
    return None


def write_csv(path, headers, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)


def run_analysis_files(connection):
    for name in (
        "04_basic_analysis.sql",
        "05_plugin_analysis.sql",
        "06_compatibility_analysis.sql",
        "07_bug_analysis.sql",
        "08_management_summary.sql",
    ):
        path = ROOT / "sql" / name
        for statement in split_statements(path.read_text(encoding="utf-8")):
            if not statement_body(statement):
                continue
            try:
                cursor = connection.execute(statement)
            except sqlite3.Error as error:
                preview = statement_body(statement)[:180]
                raise RuntimeError(f"{name} failed: {error}\n{preview}") from error
            target = export_target(statement)
            if cursor.description and target:
                headers = [column[0] for column in cursor.description]
                rows = cursor.fetchall()
                write_csv(ROOT / target, headers, rows)
                print(f"Wrote {target}")
            elif cursor.description:
                cursor.fetchall()


def assert_database_is_clean(connection):
    checks = [
        (
            "SELECT COUNT(*) FROM qa_test_results WHERE operating_system = 'Windows11'",
            "Operating system values still need cleaning.",
        ),
        (
            "SELECT COUNT(*) FROM qa_test_results WHERE daw != TRIM(daw) OR daw = 'reaper'",
            "DAW names still need cleaning.",
        ),
        (
            "SELECT COUNT(*) FROM qa_test_results WHERE test_status NOT IN ('Pass', 'Fail')",
            "Test status values still need cleaning.",
        ),
        (
            "SELECT COUNT(*) FROM qa_test_results WHERE plugin_format NOT IN ('VST3', 'AU', 'AAX')",
            "Plugin format values still need cleaning.",
        ),
        (
            "SELECT COUNT(*) FROM bugs WHERE severity NOT IN ('Low', 'Medium', 'High', 'Critical')",
            "Severity values still need cleaning.",
        ),
        (
            "SELECT COUNT(*) FROM bugs WHERE bug_status NOT IN ('Open', 'In Progress', 'Resolved')",
            "Bug status values still need cleaning.",
        ),
        (
            """
            SELECT COUNT(*) FROM bugs
            WHERE bug_status = 'Open'
              AND (date_resolved IS NOT NULL OR resolution_days IS NOT NULL)
            """,
            "An open bug still has a resolution date.",
        ),
        (
            """
            SELECT COUNT(*) FROM bugs
            WHERE bug_status = 'Resolved' AND resolution_days IS NULL
            """,
            "A resolved bug is missing resolution days.",
        ),
        (
            """
            SELECT COUNT(*) FROM (
                SELECT 1
                FROM qa_test_results
                GROUP BY plugin_id, test_date, operating_system, daw, plugin_format,
                         test_category, test_status, test_duration_minutes
                HAVING COUNT(*) > 1
            )
            """,
            "The duplicate QA test is still in the database.",
        ),
    ]
    problems = []
    for sql, message in checks:
        if connection.execute(sql).fetchone()[0]:
            problems.append(message)
    if problems:
        print("The database still has data-quality issues:")
        for problem in problems:
            print(f"- {problem}")
        print("Run sql/03_data_cleaning.sql, then run this script again.")
        sys.exit(1)


def fetch_all(connection, sql):
    cursor = connection.execute(sql)
    return cursor.fetchall()


def style_title(ws, text, subtitle):
    ws["A1"] = text
    ws["A1"].font = title_font
    ws["A1"].alignment = left_align
    ws.merge_cells("A1:F1")
    ws["A2"] = subtitle
    ws["A2"].font = note_font
    ws["A2"].alignment = left_align
    ws.merge_cells("A2:F2")
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 32
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = "Cedarline Audio QA  |  Synthetic portfolio data"


def write_headers(ws, row, headers, fill=header_fill):
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(row, column, header)
        cell.font = header_font
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin
    ws.row_dimensions[row].height = 30


def paint_body(ws, start_row, end_row, columns, percent_columns=None, number_columns=None):
    percent_columns = percent_columns or set()
    number_columns = number_columns or set()
    for row in range(start_row, end_row + 1):
        for column in range(1, columns + 1):
            cell = ws.cell(row, column)
            cell.font = label_font
            cell.border = thin
            cell.alignment = center if column > 1 else Alignment(vertical="center")
            if row % 2 == 0:
                cell.fill = pale_fill
            if column in percent_columns:
                cell.number_format = "0.0%"
            elif column in number_columns:
                cell.number_format = "#,##0.0"
        ws.row_dimensions[row].height = 20


def add_total_row(ws, row, columns, label="Total"):
    for column in range(1, columns + 1):
        cell = ws.cell(row, column)
        cell.font = total_font
        cell.fill = total_fill
        cell.border = thin
        cell.alignment = center if column > 1 else Alignment(vertical="center")
    ws.cell(row, 1, label)


def add_table(ws, name, ref):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def column_chart(title, anchor_ws, data_ref, cats_ref, anchor, width=14, height=7):
    chart = BarChart()
    chart.type = "col"
    chart.title = title
    chart.y_axis.title = None
    chart.style = 10
    chart.legend = None
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    chart.shape = 4
    chart.width = width
    chart.height = height
    chart.y_axis.majorGridlines = None
    anchor_ws.add_chart(chart, anchor)
    return chart


def rate_chart(title, ws, data_ref, cats_ref, anchor, horizontal=False):
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.title = title
    chart.style = 10
    chart.legend = None
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    chart.shape = 4
    chart.width = 16 if horizontal else 14
    chart.height = 8
    value_axis = chart.x_axis if horizontal else chart.y_axis
    value_axis.scaling.min = 0
    value_axis.numFmt = "0%"
    value_axis.title = "Failure rate"
    ws.add_chart(chart, anchor)
    return chart


def build_management_sheet(wb):
    ws = wb.active
    ws.title = "Management Summary"
    style_title(
        ws,
        "Cedarline Audio — QA Management Summary",
        "October 2025 through March 2026. Synthetic data for a portfolio project. "
        "These KPIs are Excel formulas based on the QA Test Data and Bug Data sheets.",
    )

    write_headers(ws, 4, ["KPI", "Value"])
    kpis = [
        (5, "Total QA Tests", '=COUNTA(QATests[Test ID])', "#,##0"),
        (6, "Passed Tests", '=COUNTIF(QATests[Test Status],"Pass")', "#,##0"),
        (7, "Failed Tests", '=COUNTIF(QATests[Test Status],"Fail")', "#,##0"),
        (8, "Pass Rate", "=IF(B5=0,\"\",B6/B5)", "0.0%"),
        (9, "Total Bugs", "=COUNTA(BugData[Bug ID])", "#,##0"),
        (10, "Open Bugs", '=COUNTIF(BugData[Bug Status],"Open")', "#,##0"),
        (11, "In Progress Bugs", '=COUNTIF(BugData[Bug Status],"In Progress")', "#,##0"),
        (12, "Resolved Bugs", '=COUNTIF(BugData[Bug Status],"Resolved")', "#,##0"),
        (13, "High Severity Bugs", '=COUNTIF(BugData[Severity],"High")', "#,##0"),
        (14, "Critical Bugs", '=COUNTIF(BugData[Severity],"Critical")', "#,##0"),
        (15, "Average Resolution Days", "=AVERAGE(BugData[Resolution Days])", "0.0"),
    ]
    for row, label, formula, number_format in kpis:
        ws.cell(row, 1, label).font = label_font
        ws.cell(row, 1).alignment = Alignment(vertical="center")
        ws.cell(row, 1).border = thin
        ws.cell(row, 1).fill = light_fill if row % 2 else pale_fill
        value = ws.cell(row, 2, formula)
        value.font = Font(name="Calibri", size=12, bold=True, color=NAVY)
        value.number_format = number_format
        value.alignment = center
        value.border = thin
        value.fill = light_fill if row % 2 else pale_fill
        ws.row_dimensions[row].height = 20

    ws["A17"] = "Additional checks"
    ws["A17"].font = section_font
    extra = [
        (18, "Total Testing Minutes", "=SUM(QATests[Duration Minutes])", "#,##0"),
        (19, "Average Test Duration (minutes)", "=AVERAGE(QATests[Duration Minutes])", "0.0"),
    ]
    for row, label, formula, number_format in extra:
        ws.cell(row, 1, label).font = label_font
        ws.cell(row, 1).border = thin
        ws.cell(row, 1).fill = pale_fill
        value = ws.cell(row, 2, formula)
        value.font = label_font
        value.number_format = number_format
        value.alignment = center
        value.border = thin
        value.fill = pale_fill

    ws["A21"] = (
        "Pass Rate divides Passed Tests by Total QA Tests. "
        "Average Resolution Days uses AVERAGE, which skips blank cells, so open and "
        "in-progress bugs are not included. Total Testing Minutes uses SUM."
    )
    ws["A21"].font = note_font
    ws["A21"].alignment = left_align
    ws.merge_cells("A21:F21")
    ws.row_dimensions[21].height = 36

    ws["A23"] = "Chart data"
    ws["A23"].font = section_font
    write_headers(ws, 24, ["Result", "Tests"])
    ws["A25"] = "Pass"
    ws["B25"] = "=B6"
    ws["A26"] = "Fail"
    ws["B26"] = "=B7"
    write_headers(ws, 28, ["Status", "Bugs"])
    ws["A29"] = "Open"
    ws["B29"] = "=B10"
    ws["A30"] = "In Progress"
    ws["B30"] = "=B11"
    ws["A31"] = "Resolved"
    ws["B31"] = "=B12"
    write_headers(ws, 33, ["Severity", "Bugs"])
    for row, severity in enumerate(["Low", "Medium", "High", "Critical"], start=34):
        ws.cell(row, 1, severity)
        ws.cell(row, 2, f'=COUNTIF(BugData[Severity],A{row})')
    for row in (25, 26, 29, 30, 31, 34, 35, 36, 37):
        ws.cell(row, 1).font = label_font
        ws.cell(row, 1).border = thin
        ws.cell(row, 2).font = label_font
        ws.cell(row, 2).border = thin
        ws.cell(row, 2).alignment = center
        ws.cell(row, 2).number_format = "#,##0"
        ws.cell(row, 2).fill = pale_fill

    pie = PieChart()
    pie.title = "Pass vs Fail"
    pie.add_data(Reference(ws, min_col=2, min_row=24, max_row=26), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=1, min_row=25, max_row=26))
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showVal = False
    pie.dataLabels.showCatName = False
    pie.width = 12
    pie.height = 7
    ws.add_chart(pie, "H4")

    status_chart = column_chart(
        "Bugs by Status",
        ws,
        Reference(ws, min_col=2, min_row=28, max_row=31),
        Reference(ws, min_col=1, min_row=29, max_row=31),
        "H18",
    )
    status_chart.y_axis.title = "Bugs"
    severity_chart = column_chart(
        "Bugs by Severity",
        ws,
        Reference(ws, min_col=2, min_row=33, max_row=37),
        Reference(ws, min_col=1, min_row=34, max_row=37),
        "H32",
    )
    severity_chart.y_axis.title = "Bugs"

    ws["A39"] = (
        "This workbook uses synthetic data created for educational and portfolio purposes. "
        "It does not contain proprietary data from any real audio software company."
    )
    ws["A39"].font = note_font
    ws["A39"].alignment = left_align
    ws.merge_cells("A39:F39")
    ws.row_dimensions[39].height = 32

    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 3
    ws.freeze_panes = "A4"
    ws.sheet_properties.tabColor = NAVY
    ws.oddFooter.left.text = "Synthetic portfolio data"
    ws.oddFooter.right.text = "Page &P"


def build_plugin_sheet(wb, plugin_names):
    ws = wb.create_sheet("Plugin Performance")
    style_title(
        ws,
        "Plugin Performance",
        "Failure rate = failed tests / total tests. Read the rate next to Total Tests. "
        "A high rate on a small number of tests is a weaker signal.",
    )
    headers = [
        "Plugin",
        "Total Tests",
        "Failed Tests",
        "Failure Rate",
        "Total Bugs",
        "Open Bugs",
    ]
    write_headers(ws, 4, headers)
    start = 5
    for offset, name in enumerate(plugin_names):
        row = start + offset
        ws.cell(row, 1, name)
        ws.cell(row, 2, f"=COUNTIF(QATests[Plugin],A{row})")
        ws.cell(row, 3, f'=COUNTIFS(QATests[Plugin],A{row},QATests[Test Status],"Fail")')
        ws.cell(row, 4, f'=IF(B{row}=0,"",C{row}/B{row})')
        ws.cell(row, 5, f"=COUNTIF(BugData[Plugin],A{row})")
        ws.cell(row, 6, f'=COUNTIFS(BugData[Plugin],A{row},BugData[Bug Status],"Open")')
    end = start + len(plugin_names) - 1
    paint_body(ws, start, end, 6, percent_columns={4})
    for row in range(start, end + 1):
        for column in (2, 3, 5, 6):
            ws.cell(row, column).number_format = "#,##0"

    total_row = end + 1
    add_total_row(ws, total_row, 6)
    ws.cell(total_row, 2, f"=SUM(B{start}:B{end})")
    ws.cell(total_row, 3, f"=SUM(C{start}:C{end})")
    ws.cell(total_row, 4, f'=IF(B{total_row}=0,"",C{total_row}/B{total_row})')
    ws.cell(total_row, 5, f"=SUM(E{start}:E{end})")
    ws.cell(total_row, 6, f"=SUM(F{start}:F{end})")
    ws.cell(total_row, 2).number_format = "#,##0"
    ws.cell(total_row, 3).number_format = "#,##0"
    ws.cell(total_row, 4).number_format = "0.0%"
    ws.cell(total_row, 5).number_format = "#,##0"
    ws.cell(total_row, 6).number_format = "#,##0"

    ws.conditional_formatting.add(
        f"D{start}:D{end}",
        CellIsRule(operator="greaterThan", formula=["0.2"], fill=alert_fill),
    )
    add_table(ws, "PluginSummary", f"A4:F{end}")

    rate_chart(
        "Failure Rate by Plugin",
        ws,
        Reference(ws, min_col=4, min_row=4, max_row=end),
        Reference(ws, min_col=1, min_row=start, max_row=end),
        "H4",
        horizontal=True,
    )
    note_row = total_row + 2
    ws.cell(
        note_row,
        1,
        "Rows are sorted by failure rate. The orange highlight marks a rate above 20%. "
        "The Total row uses SUM. It sits outside the filter table so the chart does not include it.",
    )
    ws.cell(note_row, 1).font = note_font
    ws.cell(note_row, 1).alignment = left_align
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=6)
    ws.row_dimensions[note_row].height = 32

    widths = [36, 14, 14, 14, 14, 16]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A5"
    ws.sheet_properties.tabColor = "2F6F8F"


def build_compatibility_sheet(wb, daws, formats, combos):
    ws = wb.create_sheet("Compatibility")
    style_title(
        ws,
        "Compatibility",
        "Compare failure rate with the number of tests. "
        "Logic Pro is tested as AU. Pro Tools is tested as AAX. "
        "Ableton Live and REAPER are tested mainly as VST3.",
    )
    ws["A4"] = "By DAW"
    ws["A4"].font = section_font
    write_headers(ws, 5, ["DAW", "Total Tests", "Failed Tests", "Failure Rate"])
    daw_start = 6
    for offset, daw in enumerate(daws):
        row = daw_start + offset
        ws.cell(row, 1, daw)
        ws.cell(row, 2, f"=COUNTIF(QATests[DAW],A{row})")
        ws.cell(row, 3, f'=COUNTIFS(QATests[DAW],A{row},QATests[Test Status],"Fail")')
        ws.cell(row, 4, f'=IF(B{row}=0,"",C{row}/B{row})')
    daw_end = daw_start + len(daws) - 1
    paint_body(ws, daw_start, daw_end, 4, percent_columns={4})
    for row in range(daw_start, daw_end + 1):
        ws.cell(row, 2).number_format = "#,##0"
        ws.cell(row, 3).number_format = "#,##0"
    ws.conditional_formatting.add(
        f"D{daw_start}:D{daw_end}",
        CellIsRule(operator="greaterThan", formula=["0.2"], fill=alert_fill),
    )

    format_header = daw_end + 2
    ws.cell(format_header, 1, "By plugin format").font = section_font
    write_headers(
        ws,
        format_header + 1,
        ["Plugin Format", "Total Tests", "Failed Tests", "Failure Rate"],
    )
    format_start = format_header + 2
    for offset, plugin_format in enumerate(formats):
        row = format_start + offset
        ws.cell(row, 1, plugin_format)
        ws.cell(row, 2, f"=COUNTIF(QATests[Plugin Format],A{row})")
        ws.cell(row, 3, f'=COUNTIFS(QATests[Plugin Format],A{row},QATests[Test Status],"Fail")')
        ws.cell(row, 4, f'=IF(B{row}=0,"",C{row}/B{row})')
    format_end = format_start + len(formats) - 1
    paint_body(ws, format_start, format_end, 4, percent_columns={4})
    for row in range(format_start, format_end + 1):
        ws.cell(row, 2).number_format = "#,##0"
        ws.cell(row, 3).number_format = "#,##0"

    combo_header = format_end + 2
    ws.cell(combo_header, 1, "By DAW and plugin format").font = section_font
    ws.cell(combo_header + 1, 1, (
        "COUNTIFS calculates each combination. A high rate on a small number of tests "
        "is a weak signal. The SQL report uses the same rule and keeps groups with at least 5 tests."
    ))
    ws.cell(combo_header + 1, 1).font = note_font
    ws.cell(combo_header + 1, 1).alignment = left_align
    ws.merge_cells(start_row=combo_header + 1, start_column=1, end_row=combo_header + 1, end_column=5)
    ws.row_dimensions[combo_header + 1].height = 32
    write_headers(
        ws,
        combo_header + 2,
        ["DAW", "Plugin Format", "Total Tests", "Failed Tests", "Failure Rate"],
    )
    combo_start = combo_header + 3
    for offset, (daw, plugin_format) in enumerate(combos):
        row = combo_start + offset
        ws.cell(row, 1, daw)
        ws.cell(row, 2, plugin_format)
        ws.cell(
            row,
            3,
            f'=COUNTIFS(QATests[DAW],A{row},QATests[Plugin Format],B{row})',
        )
        ws.cell(
            row,
            4,
            f'=COUNTIFS(QATests[DAW],A{row},QATests[Plugin Format],B{row},QATests[Test Status],"Fail")',
        )
        ws.cell(row, 5, f'=IF(C{row}=0,"",D{row}/C{row})')
    combo_end = combo_start + len(combos) - 1
    paint_body(ws, combo_start, combo_end, 5, percent_columns={5})
    for row in range(combo_start, combo_end + 1):
        ws.cell(row, 3).number_format = "#,##0"
        ws.cell(row, 4).number_format = "#,##0"
    ws.conditional_formatting.add(
        f"E{combo_start}:E{combo_end}",
        CellIsRule(operator="greaterThan", formula=["0.2"], fill=alert_fill),
    )

    rate_chart(
        "Failure Rate by DAW",
        ws,
        Reference(ws, min_col=4, min_row=5, max_row=daw_end),
        Reference(ws, min_col=1, min_row=daw_start, max_row=daw_end),
        "G4",
        horizontal=False,
    )
    for index, width in enumerate([22, 18, 16, 16, 16], start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A6"
    ws.sheet_properties.tabColor = "3D6B4F"


def build_bug_sheet(wb, categories):
    ws = wb.create_sheet("Bug Analysis")
    style_title(
        ws,
        "Bug Analysis",
        "GUI issues are counted separately from crashes. Average days use AVERAGEIF "
        "and ignore blank resolution days, so only resolved bugs are included.",
    )
    ws["A4"] = "By category"
    ws["A4"].font = section_font
    write_headers(ws, 5, ["Bug Category", "Bug Count"])
    start = 6
    for offset, category in enumerate(categories):
        row = start + offset
        ws.cell(row, 1, category)
        ws.cell(row, 2, f"=COUNTIF(BugData[Bug Category],A{row})")
    end = start + len(categories) - 1
    paint_body(ws, start, end, 2)
    for row in range(start, end + 1):
        ws.cell(row, 2).number_format = "#,##0"
    total_row = end + 1
    add_total_row(ws, total_row, 2)
    ws.cell(total_row, 2, f"=SUM(B{start}:B{end})")
    ws.cell(total_row, 2).number_format = "#,##0"

    severity_title = total_row + 2
    ws.cell(severity_title, 1, "By severity").font = section_font
    write_headers(
        ws,
        severity_title + 1,
        ["Severity", "Bug Count", "Average Resolution Days"],
    )
    severity_start = severity_title + 2
    for offset, severity in enumerate(["Low", "Medium", "High", "Critical"]):
        row = severity_start + offset
        ws.cell(row, 1, severity)
        ws.cell(row, 2, f"=COUNTIF(BugData[Severity],A{row})")
        ws.cell(
            row,
            3,
            f'=IFERROR(AVERAGEIF(BugData[Severity],A{row},BugData[Resolution Days]),"")',
        )
    severity_end = severity_start + 3
    paint_body(ws, severity_start, severity_end, 3, number_columns={3})
    for row in range(severity_start, severity_end + 1):
        ws.cell(row, 2).number_format = "#,##0"

    status_title = severity_end + 2
    ws.cell(status_title, 1, "By status").font = section_font
    write_headers(ws, status_title + 1, ["Bug Status", "Bug Count"])
    status_start = status_title + 2
    for offset, status in enumerate(["Open", "In Progress", "Resolved"]):
        row = status_start + offset
        ws.cell(row, 1, status)
        ws.cell(row, 2, f"=COUNTIF(BugData[Bug Status],A{row})")
    status_end = status_start + 2
    paint_body(ws, status_start, status_end, 2)
    for row in range(status_start, status_end + 1):
        ws.cell(row, 2).number_format = "#,##0"

    category_chart = BarChart()
    category_chart.type = "bar"
    category_chart.title = "Bugs by Category"
    category_chart.style = 10
    category_chart.legend = None
    category_chart.add_data(Reference(ws, min_col=2, min_row=5, max_row=end), titles_from_data=True)
    category_chart.set_categories(Reference(ws, min_col=1, min_row=start, max_row=end))
    category_chart.width = 16
    category_chart.height = 8
    category_chart.x_axis.title = "Bugs"
    ws.add_chart(category_chart, "E4")

    severity_chart = column_chart(
        "Bugs by Severity",
        ws,
        Reference(ws, min_col=2, min_row=severity_title + 1, max_row=severity_end),
        Reference(ws, min_col=1, min_row=severity_start, max_row=severity_end),
        "E22",
    )
    severity_chart.y_axis.title = "Bugs"

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 28
    ws.freeze_panes = "A6"
    ws.sheet_properties.tabColor = "A15C38"


def write_data_sheet(ws, title, note, headers, rows, table_name, date_columns, tab_color):
    ws["A1"] = title
    ws["A1"].font = title_font
    ws.merge_cells("A1:D1")
    ws["A2"] = note
    ws["A2"].font = note_font
    ws["A2"].alignment = left_align
    ws.merge_cells("A2:G2")
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 32
    header_row = 4
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(header_row, column, header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
    for row_index, record in enumerate(rows, start=header_row + 1):
        for column, header in enumerate(headers, start=1):
            value = record[header]
            cell = ws.cell(row_index, column, value)
            cell.font = label_font
            if header in date_columns and value:
                cell.value = datetime.strptime(value, "%Y-%m-%d").date()
                cell.number_format = "YYYY-MM-DD"
            elif isinstance(value, int):
                cell.number_format = "0"
    last_row = header_row + len(rows)
    last_column = get_column_letter(len(headers))
    add_table(ws, table_name, f"A{header_row}:{last_column}{last_row}")
    for index, header in enumerate(headers, start=1):
        width = 22 if header in ("Plugin", "Bug Category", "Test Category") else 16
        if header in ("Test ID", "Bug ID"):
            width = 12
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A5"
    ws.sheet_properties.tabColor = tab_color
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.left.text = "Synthetic portfolio data"


def build_workbook(connection):
    plugins = [
        row[0]
        for row in fetch_all(
            connection,
            """
            SELECT p.plugin_name || ' (' || p.plugin_version || ')' AS plugin_name
            FROM plugins AS p
            INNER JOIN qa_test_results AS t ON p.plugin_id = t.plugin_id
            GROUP BY p.plugin_id
            ORDER BY
                1.0 * SUM(CASE WHEN t.test_status = 'Fail' THEN 1 ELSE 0 END) / COUNT(*) DESC,
                SUM(CASE WHEN t.test_status = 'Fail' THEN 1 ELSE 0 END) DESC
            """,
        )
    ]
    daws = [
        row[0]
        for row in fetch_all(
            connection,
            """
            SELECT daw
            FROM qa_test_results
            GROUP BY daw
            ORDER BY 1.0 * SUM(CASE WHEN test_status = 'Fail' THEN 1 ELSE 0 END) / COUNT(*) DESC
            """,
        )
    ]
    formats = [
        row[0]
        for row in fetch_all(
            connection,
            """
            SELECT plugin_format
            FROM qa_test_results
            GROUP BY plugin_format
            ORDER BY 1.0 * SUM(CASE WHEN test_status = 'Fail' THEN 1 ELSE 0 END) / COUNT(*) DESC
            """,
        )
    ]
    combos = fetch_all(
        connection,
        """
        SELECT daw, plugin_format
        FROM qa_test_results
        GROUP BY daw, plugin_format
        HAVING COUNT(*) >= 5
        ORDER BY 1.0 * SUM(CASE WHEN test_status = 'Fail' THEN 1 ELSE 0 END) / COUNT(*) DESC
        """,
    )
    categories = [
        row[0]
        for row in fetch_all(
            connection,
            """
            SELECT bug_category
            FROM bugs
            GROUP BY bug_category
            ORDER BY COUNT(*) DESC, bug_category
            """,
        )
    ]
    test_headers = [
        "Test ID",
        "Plugin",
        "Test Date",
        "Operating System",
        "DAW",
        "Plugin Format",
        "Test Category",
        "Test Status",
        "Duration Minutes",
    ]
    tests = []
    for row in fetch_all(
        connection,
        """
        SELECT
            t.test_id,
            p.plugin_name || ' (' || p.plugin_version || ')' AS plugin_name,
            t.test_date,
            t.operating_system,
            t.daw,
            t.plugin_format,
            t.test_category,
            t.test_status,
            t.test_duration_minutes
        FROM qa_test_results AS t
        INNER JOIN plugins AS p ON t.plugin_id = p.plugin_id
        ORDER BY t.test_date, t.test_id
        """,
    ):
        tests.append(dict(zip(test_headers, row)))

    bug_headers = [
        "Bug ID",
        "Test ID",
        "Plugin",
        "Bug Category",
        "Severity",
        "Bug Status",
        "Date Reported",
        "Date Resolved",
        "Resolution Days",
    ]
    bugs = []
    for row in fetch_all(
        connection,
        """
        SELECT
            b.bug_id,
            b.test_id,
            p.plugin_name || ' (' || p.plugin_version || ')' AS plugin_name,
            b.bug_category,
            b.severity,
            b.bug_status,
            b.date_reported,
            b.date_resolved,
            b.resolution_days
        FROM bugs AS b
        INNER JOIN plugins AS p ON b.plugin_id = p.plugin_id
        ORDER BY b.date_reported, b.bug_id
        """,
    ):
        bugs.append(dict(zip(bug_headers, row)))

    workbook = Workbook()
    build_management_sheet(workbook)
    build_plugin_sheet(workbook, plugins)
    build_compatibility_sheet(workbook, daws, formats, combos)
    build_bug_sheet(workbook, categories)

    qa_sheet = workbook.create_sheet("QA Test Data")
    write_data_sheet(
        qa_sheet,
        "QA Test Data",
        "Cleaned test results. This is an Excel Table: use the filter arrows to sort and filter. "
        "Five tests have a blank duration. Those blanks were left blank.",
        test_headers,
        tests,
        "QATests",
        {"Test Date"},
        "5B6770",
    )
    bug_sheet = workbook.create_sheet("Bug Data")
    write_data_sheet(
        bug_sheet,
        "Bug Data",
        "Cleaned bug records used by the formulas on the summary sheets. "
        "Resolution Days is blank when a bug is Open or In Progress.",
        bug_headers,
        bugs,
        "BugData",
        {"Date Reported", "Date Resolved"},
        "5B6770",
    )

    workbook.properties.title = "Audio Plugin QA Analysis"
    workbook.properties.subject = "Synthetic portfolio QA analysis"
    workbook.properties.creator = "Cedarline Audio QA portfolio project"
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True

    EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(EXCEL_PATH)
    print(f"Wrote {EXCEL_PATH.relative_to(ROOT)}")


def main():
    if not DB_PATH.exists():
        print("Database not found. Run scripts/load_database.py first.")
        sys.exit(1)

    connection = sqlite3.connect(DB_PATH)
    try:
        connection.execute("PRAGMA foreign_keys = ON")
        assert_database_is_clean(connection)
        run_analysis_files(connection)
        build_workbook(connection)
    finally:
        connection.close()


if __name__ == "__main__":
    main()
